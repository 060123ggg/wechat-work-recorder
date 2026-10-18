from datetime import date
from decimal import Decimal
from openpyxl import load_workbook
import pytest
from wechat_work_recorder.models import WorkKind, WorkRecord
from wechat_work_recorder.settings import directory_from_settings, load_settings, save_settings
from wechat_work_recorder.parser import parse_messages
from wechat_work_recorder.workbook import write_mock_workbook


def test_settings_parse_and_confirmed_mock_workbook(tmp_path):
    settings = tmp_path / "settings.json"
    save_settings(["张三", "李四"], ["一群", "二群", "三群"], settings)
    employees, groups = load_settings(settings)
    result = parse_messages("2026-10-18 张三 货号 WX-123/ABC 数量 4 单价 1.50", directory_from_settings(employees, groups), "一群")
    assert not result.warnings
    path = tmp_path / "mock.xlsx"
    with pytest.raises(PermissionError):
        write_mock_workbook(path, result.records)
    write_mock_workbook(path, result.records, confirmed=True)
    ws = load_workbook(path, read_only=True).active
    assert ws["D2"].value == "WX-123/ABC"
    assert ws["G2"].value == 6
    assert ws["H2"].value == 6
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        write_mock_workbook(path, result.records, confirmed=True)
    assert path.read_bytes() == original


def test_settings_limits(tmp_path):
    with pytest.raises(ValueError):
        save_settings([str(i) for i in range(18)], ["一群"], tmp_path / "settings.json")


def test_failed_mock_save_removes_only_new_partial_file(tmp_path, monkeypatch):
    from openpyxl import Workbook
    record = WorkRecord("1", date(2026, 10, 18), WorkKind.HOURLY, "一群", "计时2", hours=Decimal("2"))
    target = tmp_path / "partial.xlsx"

    def fail_after_write(self, output):
        output.write(b"partial")
        raise OSError("disk failure")

    monkeypatch.setattr(Workbook, "save", fail_after_write)
    with pytest.raises(OSError):
        write_mock_workbook(target, [record], confirmed=True)
    assert not target.exists()

    target.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        write_mock_workbook(target, [record], confirmed=True)
    assert target.read_bytes() == b"existing"
